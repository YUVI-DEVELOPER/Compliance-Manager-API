from __future__ import annotations

from html import escape
from pathlib import Path
from urllib.parse import unquote, urlparse

from fastapi import HTTPException, status
from fastapi.responses import FileResponse, HTMLResponse, Response

from app.core.config import get_settings

TEXT_EXTENSIONS = {".csv", ".json", ".md", ".rtf", ".txt", ".xml"}
IMAGE_EXTENSIONS = {".gif", ".jpeg", ".jpg", ".png", ".webp"}
PDF_EXTENSIONS = {".pdf"}
DOCX_EXTENSIONS = {".docx", ".docm"}


class DocumentViewerError(HTTPException):
    def __init__(self, detail: str, status_code: int = status.HTTP_400_BAD_REQUEST) -> None:
        super().__init__(status_code=status_code, detail=detail)


def _resolve_local_document_path(source_url: str) -> Path:
    normalized = source_url.strip()
    if not normalized:
        raise DocumentViewerError("source_url is required")

    parsed = urlparse(normalized)
    path = unquote(parsed.path if parsed.scheme in {"http", "https"} else normalized)
    path = path.replace("\\", "/")

    settings = get_settings()
    candidates: list[tuple[str, Path]] = [
        ("/uploads/", Path(settings.FILE_UPLOAD_DIR).resolve()),
        ("/filestorage/", Path(settings.FILE_STORAGE_DIR).resolve()),
        ("uploads/", Path(settings.FILE_UPLOAD_DIR).resolve()),
        ("filestorage/", Path(settings.FILE_STORAGE_DIR).resolve()),
    ]

    for marker, root in candidates:
        marker_index = path.find(marker)
        if marker_index < 0:
            continue
        relative = path[marker_index + len(marker):].lstrip("/")
        resolved = (root / relative).resolve()
        try:
            resolved.relative_to(root)
        except ValueError:
            raise DocumentViewerError("Document path is outside allowed storage", status.HTTP_403_FORBIDDEN) from None
        if resolved.is_file():
            return resolved

    local = Path(normalized)
    if local.is_file():
        resolved_local = local.resolve()
        allowed_roots = [Path(settings.FILE_UPLOAD_DIR).resolve(), Path(settings.FILE_STORAGE_DIR).resolve()]
        if any(resolved_local.is_relative_to(root) for root in allowed_roots):
            return resolved_local

    raise DocumentViewerError("Only locally stored application documents can be previewed internally.", status.HTTP_404_NOT_FOUND)


def _plain_text_response(path: Path) -> Response:
    text = path.read_text(encoding="utf-8", errors="replace")
    return Response(
        content=text,
        media_type="text/plain; charset=utf-8",
        headers={"Content-Disposition": f'inline; filename="{path.name}"', "Cache-Control": "no-store"},
    )


def _docx_html_response(path: Path) -> HTMLResponse:
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:
        raise DocumentViewerError("DOCX preview requires python-docx.", status.HTTP_503_SERVICE_UNAVAILABLE) from exc

    try:
        document = DocxDocument(str(path))
    except Exception as exc:
        raise DocumentViewerError("DOCX document could not be opened for preview.") from exc

    body: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
        if "heading 1" in style_name:
            body.append(f"<h1>{escape(text)}</h1>")
        elif "heading 2" in style_name:
            body.append(f"<h2>{escape(text)}</h2>")
        elif "heading" in style_name:
            body.append(f"<h3>{escape(text)}</h3>")
        else:
            body.append(f"<p>{escape(text)}</p>")

    for table in document.tables:
        rows = []
        for row in table.rows:
            rows.append("<tr>" + "".join(f"<td>{escape(cell.text.strip())}</td>" for cell in row.cells) + "</tr>")
        if rows:
            body.append("<table>" + "".join(rows) + "</table>")

    html = f"""<!doctype html>
<html>
<head>
  <meta charset="utf-8" />
  <title>{escape(path.name)}</title>
  <style>
    body {{ color: #111827; font-family: Arial, sans-serif; margin: 36px auto; max-width: 860px; }}
    h1 {{ font-size: 24px; margin: 0 0 18px; }}
    h2 {{ border-bottom: 1px solid #d1d5db; font-size: 18px; margin-top: 26px; padding-bottom: 6px; }}
    h3 {{ font-size: 15px; margin-top: 20px; }}
    p {{ font-size: 13px; line-height: 1.55; }}
    table {{ border-collapse: collapse; margin: 16px 0; width: 100%; }}
    td {{ border: 1px solid #d1d5db; font-size: 12px; padding: 7px; vertical-align: top; }}
  </style>
</head>
<body>
  {''.join(body) if body else '<p>No previewable DOCX text was found.</p>'}
</body>
</html>"""
    return HTMLResponse(content=html, headers={"Cache-Control": "no-store"})


def build_document_preview_response(source_url: str) -> Response:
    path = _resolve_local_document_path(source_url)
    suffix = path.suffix.lower()
    headers = {"Content-Disposition": f'inline; filename="{path.name}"', "Cache-Control": "no-store"}

    if suffix in PDF_EXTENSIONS:
        return FileResponse(path, media_type="application/pdf", headers=headers)
    if suffix in IMAGE_EXTENSIONS:
        return FileResponse(path, headers=headers)
    if suffix in TEXT_EXTENSIONS:
        return _plain_text_response(path)
    if suffix in DOCX_EXTENSIONS:
        return _docx_html_response(path)

    raise DocumentViewerError("This file type is not supported by the internal viewer yet.")


def build_document_original_response(source_url: str) -> FileResponse:
    path = _resolve_local_document_path(source_url)
    return FileResponse(
        path,
        headers={"Content-Disposition": f'inline; filename="{path.name}"', "Cache-Control": "no-store"},
    )


def build_storage_file_response(storage_root: Path, storage_path: str) -> FileResponse:
    root = storage_root.resolve()
    resolved = (root / storage_path).resolve()
    try:
        resolved.relative_to(root)
    except ValueError:
        raise DocumentViewerError("Document path is outside allowed storage", status.HTTP_403_FORBIDDEN) from None
    if not resolved.is_file():
        raise DocumentViewerError("Document not found", status.HTTP_404_NOT_FOUND)
    return FileResponse(
        resolved,
        headers={"Content-Disposition": f'inline; filename="{resolved.name}"', "Cache-Control": "no-store"},
    )
